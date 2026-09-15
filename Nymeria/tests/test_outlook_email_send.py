"""Outbound side of the Outlook tools: attachments, HTML replies, send_draft, DOCX.

B20 (inline attachment path and the untouched no-attachment sendMail), B21
(upload session chunking, size ceiling, secrets denylist, workspace
confinement), B22 (HTML reply and forward through the draft path, quoting
kept or dropped), B23 (outlook_send_draft), B24 (edit_draft attachment add and
remove), B25 (DOCX via python-docx, never Gemini; legacy .doc refused), B26
(save_to writes raw files into the workspace).
"""

import base64
import io

import pytest

from nymeria.tools import filesystem as fs
from nymeria.tools import outlook_attachments as oa
from nymeria.tools import outlook_email as oe
from nymeria.tools import outlook_graph as og

_CFG = {"configurable": {"user_id": "u1", "thread_id": "t1"}}


class _Graph:
    """Scripted graph_request: answers by (method, endpoint-suffix) and records calls."""

    def __init__(self, script):
        self.script = list(script)
        self.calls = []

    def __call__(self, user_id, method, endpoint, account_id=None, json_data=None, params=None, **kw):
        self.calls.append({"method": method, "endpoint": endpoint, "json": json_data, "params": params, "mailbox": kw.get("mailbox")})
        for i, (m, suffix, response) in enumerate(self.script):
            if m == method and endpoint.endswith(suffix):
                self.script.pop(i)
                return response
        raise AssertionError(f"unexpected Graph call {method} {endpoint}")

    def find(self, method, suffix):
        return [c for c in self.calls if c["method"] == method and c["endpoint"].endswith(suffix)]


def _install(monkeypatch, script):
    g = _Graph(script)
    monkeypatch.setattr(oe, "graph_request", g)
    monkeypatch.setattr(og, "graph_request", g)
    return g


@pytest.fixture
def small_file(tmp_path):
    p = tmp_path / "report.txt"
    p.write_bytes(b"hello attachment")
    return p


# ---------------------------------------------------------------------------
# B20: send with and without attachments
# ---------------------------------------------------------------------------


def test_send_without_attachments_still_uses_send_mail(monkeypatch):
    g = _install(monkeypatch, [("POST", "/me/sendMail", (True, {}))])
    out = oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B"}, config=_CFG)
    assert out == "[Success]: Email sent to a@x.com"
    assert len(g.calls) == 1
    assert g.calls[0]["json"]["message"]["toRecipients"] == [{"emailAddress": {"address": "a@x.com"}}]
    assert "importance" not in g.calls[0]["json"]["message"]


def test_send_reply_to_and_importance_ride_on_the_message(monkeypatch):
    g = _install(monkeypatch, [("POST", "/me/sendMail", (True, {}))])
    oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "reply_to": "r@x.com", "importance": "High"}, config=_CFG)
    msg = g.calls[0]["json"]["message"]
    assert msg["replyTo"] == [{"emailAddress": {"address": "r@x.com"}}]
    assert msg["importance"] == "high"
    assert oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "importance": "urgent"}, config=_CFG) == (
        "[Error]: importance must be high, normal, or low."
    )


def test_send_with_small_attachment_drafts_attaches_inline_then_sends(monkeypatch, small_file):
    g = _install(monkeypatch, [
        ("POST", "/me/messages", (True, {"id": "D1"})),
        ("POST", "/me/messages/D1/attachments", (True, {"id": "A1"})),
        ("POST", "/me/messages/D1/send", (True, {})),
    ])
    monkeypatch.setattr(oe, "require_scopes", lambda *a, **k: None)
    out = oe.outlook_send_email.invoke(
        {"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(small_file), "mailbox": "s@x.com"}, config=_CFG,
    )
    assert out == "[Success]: Email sent to a@x.com with attachments: report.txt (16 B)"
    assert [c["method"] + " " + c["endpoint"] for c in g.calls] == [
        "POST /me/messages", "POST /me/messages/D1/attachments", "POST /me/messages/D1/send",
    ]
    att = g.calls[1]["json"]
    assert att["@odata.type"] == "#microsoft.graph.fileAttachment"
    assert att["name"] == "report.txt" and att["contentType"] == "text/plain"
    assert base64.b64decode(att["contentBytes"]) == b"hello attachment"
    assert all(c["mailbox"] == "s@x.com" for c in g.calls)
    assert g.calls[0]["json"]["subject"] == "S"  # the draft IS the message object


def test_send_attach_failure_leaves_the_draft_and_does_not_send(monkeypatch, small_file):
    g = _install(monkeypatch, [
        ("POST", "/me/messages", (True, {"id": "D1"})),
        ("POST", "/me/messages/D1/attachments", (False, "API Error (413): too big")),
    ])
    out = oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(small_file)}, config=_CFG)
    assert out == "[Error]: Attaching 'report.txt' failed: API Error (413): too big. The unsent draft (ID: D1) is in Drafts."
    assert g.find("POST", "/send") == []


def test_missing_attachment_path_refuses_before_any_request(monkeypatch, tmp_path):
    g = _install(monkeypatch, [])
    out = oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(tmp_path / "nope.pdf")}, config=_CFG)
    assert out.startswith("[Error]: Attachment not found: ")
    assert g.calls == []


# ---------------------------------------------------------------------------
# B21: upload session, ceiling, denylist, confinement
# ---------------------------------------------------------------------------


def test_large_attachment_uses_an_upload_session_in_ordered_chunks(monkeypatch, tmp_path):
    monkeypatch.setattr(oa, "INLINE_ATTACHMENT_LIMIT", 100)
    monkeypatch.setattr(oa, "UPLOAD_CHUNK_BYTES", 64)
    big = tmp_path / "big.bin"
    payload = bytes(range(256)) * 1  # 256 bytes: 4 chunks of 64
    big.write_bytes(payload)
    puts = []

    def fake_put(url, content, *, content_range, timeout=120):
        puts.append((url, content, content_range))
        return True, {"status": 200 if len(puts) < 4 else 201, "headers": {}, "body": {}}

    monkeypatch.setattr(og, "graph_upload_put", fake_put)
    g = _install(monkeypatch, [
        ("POST", "/me/messages", (True, {"id": "D1"})),
        ("POST", "/me/messages/D1/attachments/createUploadSession", (True, {"uploadUrl": "https://outlook.office.com/api/v2.0/Users('u')/Messages('D1')/AttachmentSessions('S')?authtoken=x"})),
        ("POST", "/me/messages/D1/send", (True, {})),
    ])
    out = oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(big)}, config=_CFG)
    assert out == "[Success]: Email sent to a@x.com with attachments: big.bin (256 B)"
    session = g.find("POST", "createUploadSession")[0]["json"]
    assert session == {"AttachmentItem": {"attachmentType": "file", "name": "big.bin", "size": 256}}
    assert [r for _, _, r in puts] == ["bytes 0-63/256", "bytes 64-127/256", "bytes 128-191/256", "bytes 192-255/256"]
    assert b"".join(c for _, c, _ in puts) == payload
    assert g.find("POST", "/attachments") == []  # no inline POST for a large file


def test_upload_put_refuses_a_non_outlook_host_and_sends_no_auth(monkeypatch):
    ok, err = og.graph_upload_put("https://evil.example/x", b"abc", content_range="bytes 0-2/3")
    assert not ok and "Outlook service host" in err
    seen = {}

    class _C:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def request(self, method, url, **kw):
            seen.update(method=method, url=url, **kw)

            class R:
                status_code = 200
                text = ""
                headers = {}

            return R()

    monkeypatch.setattr(og, "_http_client", lambda **k: _C())
    ok, res = og.graph_upload_put("https://outlook.office.com/api/v2.0/x?authtoken=t", b"abc", content_range="bytes 0-2/3")
    assert ok and res["status"] == 200
    assert seen["method"] == "PUT" and "Authorization" not in seen["headers"]
    assert seen["headers"]["Content-Range"] == "bytes 0-2/3" and seen["headers"]["Content-Length"] == "3"


def test_oversized_attachment_is_refused_with_the_limit(monkeypatch, tmp_path):
    monkeypatch.setattr(oa, "MAX_OUTBOUND_ATTACHMENT", 10)
    p = tmp_path / "huge.bin"
    p.write_bytes(b"x" * 11)
    g = _install(monkeypatch, [])
    out = oe.outlook_create_draft.invoke({"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(p)}, config=_CFG)
    assert out == "[Error]: Attachment 'huge.bin' is 11 B; the limit is 10 B. Share a link to the file instead."
    assert g.calls == []


def test_secrets_denylist_wording_is_the_filesystem_tools(monkeypatch, small_file):
    monkeypatch.setattr(fs, "secrets_path_error", lambda path: f"Refusing to read credential store {path.name}")
    g = _install(monkeypatch, [])
    out = oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(small_file)}, config=_CFG)
    assert out == "[Error]: Refusing to read credential store report.txt"
    assert g.calls == []


def test_workspace_confinement_applies_to_outbound_attachments(monkeypatch, small_file, tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    monkeypatch.setattr(fs, "confine_file_tools_to_workspace", lambda: True)
    monkeypatch.setattr(fs, "get_workspace_dir", lambda: workspace)
    g = _install(monkeypatch, [])
    out = oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(small_file)}, config=_CFG)
    assert out.startswith("[Error]: Attachment outside workspace: ")
    assert g.calls == []
    inside = workspace / "ok.txt"
    inside.write_bytes(b"k")
    _install(monkeypatch, [("POST", "/me/messages", (True, {"id": "D"})), ("POST", "/me/messages/D/attachments", (True, {})), ("POST", "/me/messages/D/send", (True, {}))])
    assert oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B", "attachments": str(inside)}, config=_CFG).startswith("[Success]")


# ---------------------------------------------------------------------------
# B22: HTML reply and forward via the draft path
# ---------------------------------------------------------------------------

_REPLY_DRAFT = {"id": "R1", "subject": "RE: Hi", "toRecipients": [{"emailAddress": {"address": "them@y.com"}}],
                "body": {"contentType": "html", "content": "<blockquote>original</blockquote>"}}


def test_plain_reply_without_attachments_keeps_the_one_shot_action(monkeypatch):
    g = _install(monkeypatch, [("POST", "/me/messages/M1/replyAll", (True, {}))])
    out = oe.outlook_reply_email.invoke({"email_id": "M1", "body": "Thanks", "reply_all": True}, config=_CFG)
    assert out == "[Success]: Reply sent to all recipients"
    assert g.calls[0]["json"] == {"comment": "Thanks"}


def test_html_reply_quotes_the_original_beneath_and_sends(monkeypatch):
    g = _install(monkeypatch, [
        ("POST", "/me/messages/M1/createReply", (True, dict(_REPLY_DRAFT))),
        ("PATCH", "/me/messages/R1", (True, {})),
        ("POST", "/me/messages/R1/send", (True, {})),
    ])
    out = oe.outlook_reply_email.invoke({"email_id": "M1", "body": "<p>Thanks</p>", "is_html": True}, config=_CFG)
    assert out == "[Success]: Reply sent to sender"
    body = g.find("PATCH", "/R1")[0]["json"]["body"]
    assert body == {"contentType": "html", "content": "<p>Thanks</p><br><br><blockquote>original</blockquote>"}


def test_plain_reply_with_attachment_escapes_text_over_the_html_quote(monkeypatch, small_file):
    g = _install(monkeypatch, [
        ("POST", "/me/messages/M1/createReply", (True, dict(_REPLY_DRAFT))),
        ("PATCH", "/me/messages/R1", (True, {})),
        ("POST", "/me/messages/R1/attachments", (True, {})),
        ("POST", "/me/messages/R1/send", (True, {})),
    ])
    out = oe.outlook_reply_email.invoke({"email_id": "M1", "body": "1 < 2\nok", "attachments": str(small_file)}, config=_CFG)
    assert out == "[Success]: Reply sent to sender with attachments: report.txt (16 B)"
    body = g.find("PATCH", "/R1")[0]["json"]["body"]
    assert body["contentType"] == "html"
    assert body["content"].startswith('<div style="white-space:pre-wrap">1 &lt; 2\nok</div><br><br><blockquote>')


def test_quote_original_false_replaces_the_body(monkeypatch):
    g = _install(monkeypatch, [
        ("POST", "/me/messages/M1/createReply", (True, dict(_REPLY_DRAFT))),
        ("PATCH", "/me/messages/R1", (True, {})),
        ("POST", "/me/messages/R1/send", (True, {})),
    ])
    oe.outlook_reply_email.invoke({"email_id": "M1", "body": "Only this", "quote_original": False}, config=_CFG)
    assert g.find("PATCH", "/R1")[0]["json"]["body"] == {"contentType": "text", "content": "Only this"}


def test_draft_reply_keeps_the_quote_lists_attachments_and_does_not_send(monkeypatch, small_file):
    g = _install(monkeypatch, [
        ("POST", "/me/messages/M1/createReplyAll", (True, dict(_REPLY_DRAFT))),
        ("PATCH", "/me/messages/R1", (True, {})),
        ("POST", "/me/messages/R1/attachments", (True, {})),
    ])
    out = oe.outlook_draft_reply.invoke({"email_id": "M1", "body": "Hi", "reply_all": True, "attachments": str(small_file)}, config=_CFG)
    assert out == (
        "[Success]: Draft reply-all created for 'RE: Hi'\n"
        "  To: them@y.com\n"
        "  Draft ID: R1\n"
        "  Attachments: report.txt (16 B)\n"
        "  Status: In Drafts folder, ready for review and send (outlook_send_draft sends it)."
    )
    assert g.find("POST", "/send") == []
    assert "<blockquote>original</blockquote>" in g.find("PATCH", "/R1")[0]["json"]["body"]["content"]


def test_forward_plain_keeps_one_shot_and_html_goes_via_create_forward(monkeypatch):
    g = _install(monkeypatch, [("POST", "/me/messages/M1/forward", (True, {}))])
    assert oe.outlook_forward_email.invoke({"email_id": "M1", "to": "a@x.com", "comment": "FYI"}, config=_CFG) == "[Success]: Email forwarded to a@x.com"
    assert g.calls[0]["json"] == {"toRecipients": [{"emailAddress": {"address": "a@x.com"}}], "comment": "FYI"}

    g = _install(monkeypatch, [
        ("POST", "/me/messages/M1/createForward", (True, {"id": "F1", "body": {"contentType": "html", "content": "<div>orig</div>"}})),
        ("PATCH", "/me/messages/F1", (True, {})),
        ("POST", "/me/messages/F1/send", (True, {})),
    ])
    out = oe.outlook_forward_email.invoke({"email_id": "M1", "to": "a@x.com, b@x.com", "comment": "<b>FYI</b>", "is_html": True}, config=_CFG)
    assert out == "[Success]: Email forwarded to a@x.com, b@x.com"
    patch = g.find("PATCH", "/F1")[0]["json"]
    assert patch["toRecipients"] == [{"emailAddress": {"address": "a@x.com"}}, {"emailAddress": {"address": "b@x.com"}}]
    assert patch["body"] == {"contentType": "html", "content": "<b>FYI</b><br><br><div>orig</div>"}


# ---------------------------------------------------------------------------
# B23: send_draft
# ---------------------------------------------------------------------------


def test_send_draft_sends_a_draft_and_reports_recipients(monkeypatch):
    g = _install(monkeypatch, [
        ("GET", "/me/messages/D1", (True, {"id": "D1", "subject": "Quote", "isDraft": True, "hasAttachments": True,
                                          "toRecipients": [{"emailAddress": {"address": "a@x.com"}}]})),
        ("POST", "/me/messages/D1/send", (True, {})),
    ])
    out = oe.outlook_send_draft.invoke({"draft_id": "D1"}, config=_CFG)
    assert out == "[Success]: Sent 'Quote' to a@x.com (with attachments)."
    assert g.calls[0]["params"]["$select"] == "id,subject,isDraft,toRecipients,ccRecipients,hasAttachments"


def test_send_draft_refuses_non_drafts_and_empty_recipients(monkeypatch):
    g = _install(monkeypatch, [("GET", "/me/messages/S1", (True, {"id": "S1", "subject": "Sent already", "isDraft": False}))])
    assert oe.outlook_send_draft.invoke({"draft_id": "S1"}, config=_CFG) == (
        "[Error]: Message 'Sent already' is not a draft (already sent or received); nothing sent."
    )
    assert g.find("POST", "/send") == []
    g = _install(monkeypatch, [("GET", "/me/messages/D2", (True, {"id": "D2", "subject": "x", "isDraft": True, "toRecipients": []}))])
    assert oe.outlook_send_draft.invoke({"draft_id": "D2"}, config=_CFG) == (
        "[Error]: The draft has no recipients; add them with outlook_edit_draft first."
    )
    assert oe.outlook_send_draft.invoke({"draft_id": " "}, config=_CFG) == "[Error]: draft_id is required."


# ---------------------------------------------------------------------------
# B24: edit_draft attachments
# ---------------------------------------------------------------------------


def test_edit_draft_removes_by_name_and_adds_files(monkeypatch, small_file):
    g = _install(monkeypatch, [
        ("PATCH", "/me/messages/D1", (True, {})),
        ("GET", "/me/messages/D1/attachments", (True, {"value": [{"id": "A1", "name": "Old.PDF"}, {"id": "A2", "name": "keep.txt"}]})),
        ("DELETE", "/me/messages/D1/attachments/A1", (True, {})),
        ("POST", "/me/messages/D1/attachments", (True, {})),
    ])
    out = oe.outlook_edit_draft.invoke({"draft_id": "D1", "subject": "New", "attachments": f"-old.pdf, {small_file}, -ghost.txt"}, config=_CFG)
    assert out == "[Success]: Draft updated (subject, attachments removed: Old.PDF (not found: ghost.txt), attachments added: report.txt (16 B)). ID: D1"
    assert g.find("DELETE", "/A2") == []


def test_edit_draft_attachments_only_needs_no_patch(monkeypatch, small_file):
    g = _install(monkeypatch, [("POST", "/me/messages/D1/attachments", (True, {}))])
    out = oe.outlook_edit_draft.invoke({"draft_id": "D1", "attachments": str(small_file)}, config=_CFG)
    assert out == "[Success]: Draft updated (attachments added: report.txt (16 B)). ID: D1"
    assert g.find("PATCH", "/D1") == []
    assert oe.outlook_edit_draft.invoke({"draft_id": "D1"}, config=_CFG) == (
        "[Error]: No fields to update. Provide at least one of: body, subject, to, cc, bcc, attachments."
    )


# ---------------------------------------------------------------------------
# B25: DOCX via python-docx, never Gemini
# ---------------------------------------------------------------------------


def _docx_bytes():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Decision paper")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "PART-1"
    table.rows[0].cells[1].text = "qty 3"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extracts_paragraphs_and_tables_without_gemini(monkeypatch):
    monkeypatch.setattr(oa, "_extract_with_gemini", lambda *a, **k: pytest.fail("Gemini must not see a DOCX"))
    b64 = base64.b64encode(_docx_bytes()).decode()
    out = oa._extract_attachment({"name": "paper.docx", "mime_type": oa._DOCX_MIME, "data_b64": b64})
    assert "Decision paper" in out and "PART-1 | qty 3" in out
    # Extension fallback when the sender labels it octet-stream.
    out2 = oa._extract_attachment({"name": "paper.docx", "mime_type": "application/octet-stream", "data_b64": b64})
    assert "Decision paper" in out2


def test_pdf_still_routes_to_gemini_and_legacy_doc_is_refused(monkeypatch):
    seen = {}
    monkeypatch.setattr(oa, "_extract_with_gemini", lambda data, mime, name: seen.update(mime=mime, name=name) or "ok")
    assert oa._extract_attachment({"name": "q.pdf", "mime_type": "application/pdf", "data_b64": "QQ=="}) == "ok"
    assert seen == {"mime": "application/pdf", "name": "q.pdf"}
    out = oa._extract_attachment({"name": "old.doc", "mime_type": oa._LEGACY_DOC_MIME, "data_b64": "QQ=="})
    assert out.startswith("[Skipped]: 'old.doc' is a legacy binary .doc file")


# ---------------------------------------------------------------------------
# B26: save_to
# ---------------------------------------------------------------------------


def test_save_to_writes_raw_files_and_reports_paths(monkeypatch, tmp_path):
    target = tmp_path / "ws" / "attachments"
    monkeypatch.setattr(fs, "resolve_workspace_write_path", lambda p: (target, None))
    monkeypatch.setattr(oa, "_download_attachments", lambda user_id, email_id, account_id=None, mailbox=None: ([
        {"name": "notes.txt", "mime_type": "text/plain", "data_b64": base64.b64encode(b"hi").decode(), "size": 4, "skipped": False},
        {"name": "big.bin", "mime_type": "application/octet-stream", "data_b64": None, "size": 99, "skipped": True},
    ], 0))
    out = oa.outlook_get_attachments.invoke({"email_id": "M1", "save_to": "attachments"}, config=_CFG)
    saved = target / "notes.txt"
    assert saved.read_bytes() == b"hi"
    assert f"Saved:\n  - {saved}" in out
    assert not (target / "big.bin").exists()
    assert "=== Attachment 1/2: notes.txt" in out and "hi" in out


def test_save_to_outside_workspace_warns_but_extraction_still_returns(monkeypatch):
    monkeypatch.setattr(fs, "resolve_workspace_write_path", lambda p: (None, "Path outside workspace: /etc"))
    monkeypatch.setattr(oa, "_download_attachments", lambda user_id, email_id, account_id=None, mailbox=None: ([
        {"name": "notes.txt", "mime_type": "text/plain", "data_b64": base64.b64encode(b"hi").decode(), "size": 4, "skipped": False},
    ], 0))
    out = oa.outlook_get_attachments.invoke({"email_id": "M1", "save_to": "/etc"}, config=_CFG)
    assert out.startswith("[Attachment: notes.txt (text/plain)]\n\nhi")
    assert "[Warning]: save_to failed: Path outside workspace: /etc" in out


# ---------------------------------------------------------------------------
# Review fix: sending from a shared mailbox is gated on Mail.Send.Shared
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("tool,args", [
    ("outlook_send_email", {"to": "a@x.com", "subject": "S", "body": "B"}),
    ("outlook_reply_email", {"email_id": "M1", "body": "B"}),
    ("outlook_forward_email", {"email_id": "M1", "to": "a@x.com"}),
    ("outlook_send_draft", {"draft_id": "D1"}),
])
def test_shared_mailbox_send_is_refused_when_the_scope_is_missing(monkeypatch, tool, args):
    g = _install(monkeypatch, [])
    asked = []

    def _scopes(user_id, needed, account_id=None, *, mailbox=None, thread_id=None, purpose=""):
        asked.append((list(needed), mailbox, purpose))
        return "the connected account lacks Mail.Send.Shared; reconnect with request_credential(...)"

    monkeypatch.setattr(oe, "require_scopes", _scopes)
    out = getattr(oe, tool).invoke({**args, "mailbox": "sales@x.com"}, config=_CFG)
    assert out == "[Error]: the connected account lacks Mail.Send.Shared; reconnect with request_credential(...)"
    assert asked == [(["Mail.Send.Shared"], "sales@x.com", "sending from a shared mailbox")]
    assert g.calls == []  # nothing composed, nothing sent


def test_own_mailbox_send_never_consults_the_shared_scope(monkeypatch):
    g = _install(monkeypatch, [("POST", "/me/sendMail", (True, {}))])
    monkeypatch.setattr(oe, "require_scopes", lambda *a, **k: (_ for _ in ()).throw(AssertionError("gate consulted")))
    out = oe.outlook_send_email.invoke({"to": "a@x.com", "subject": "S", "body": "B"}, config=_CFG)
    assert out.startswith("[Success]")
    assert [c["endpoint"] for c in g.calls] == ["/me/sendMail"]
